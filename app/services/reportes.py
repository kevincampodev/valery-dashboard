import io
from datetime import date

from flask import current_app
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from sqlalchemy.orm import joinedload, selectinload

from ..models import Venta, Factura, Pago, Liquidacion, Meta, CANALES
from ..services.comisiones import etiqueta_quincena
from ..services.ventas import resumen_mes
from ..utils import nombre_mes

FORMATO_PESOS = '"$"#,##0;-"$"#,##0'
FORMATO_FECHA = "DD/MM/YYYY"
FORMATO_PCT = "0.0%"
FILA_ENCABEZADO = 3


def escribir_hoja(ws, titulo, encabezados, filas, formatos=None, totales=None):
    """Escribe una tabla con título, encabezado con estilo, formatos por columna y fila de totales."""
    ws.append([titulo])
    ws["A1"].font = Font(bold=True, size=14)
    ws.append([])
    ws.append(encabezados)
    for celda in ws[FILA_ENCABEZADO]:
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="343A40")

    for fila in filas:
        ws.append(list(fila))
    ultima_fila_datos = ws.max_row

    for columna, formato in (formatos or {}).items():
        for fila in range(FILA_ENCABEZADO + 1, ultima_fila_datos + 1):
            ws.cell(row=fila, column=columna).number_format = formato

    if totales and filas:
        fila_total = ultima_fila_datos + 1
        ws.cell(row=fila_total, column=1, value="TOTAL").font = Font(bold=True)
        for columna in totales:
            letra = get_column_letter(columna)
            celda = ws.cell(row=fila_total, column=columna,
                            value=f"=SUM({letra}{FILA_ENCABEZADO + 1}:{letra}{ultima_fila_datos})")
            celda.font = Font(bold=True)
            celda.number_format = FORMATO_PESOS

    for columna in range(1, len(encabezados) + 1):
        largo = max(len(str(ws.cell(row=f, column=columna).value or ""))
                    for f in range(FILA_ENCABEZADO, ultima_fila_datos + 1))
        ws.column_dimensions[get_column_letter(columna)].width = max(12, min(45, largo + 2))

    ws.freeze_panes = ws.cell(row=FILA_ENCABEZADO + 1, column=1)


def generar_reporte_mensual(inicio, fin):
    hoy = date.today()
    mes = nombre_mes(inicio)
    pesos = lambda *columnas: {c: FORMATO_PESOS for c in columnas}  # noqa: E731

    ventas = Venta.query.options(joinedload(Venta.vendedora)).filter(Venta.fecha.between(inicio, fin)).all()
    r = resumen_mes(ventas, inicio, fin)
    facturas = Factura.query.options(joinedload(Factura.proveedor), selectinload(Factura.aplicaciones)).all()
    pendientes = sorted((f for f in facturas if f.saldo > 0), key=lambda f: f.fecha_limite)
    pagos = (Pago.query.options(joinedload(Pago.proveedor))
             .filter(Pago.fecha.between(inicio, fin)).order_by(Pago.fecha).all())
    liquidaciones = (Liquidacion.query.options(joinedload(Liquidacion.vendedora))
                     .filter(Liquidacion.inicio.between(inicio, fin)).order_by(Liquidacion.inicio).all())
    meta = Meta.query.filter_by(periodo=inicio, canal="Total").first()

    wb = Workbook()

    # 1. Resumen
    ws = wb.active
    ws.title = "Resumen"
    escribir_hoja(ws, f"{current_app.config['EMPRESA_NOMBRE']} — Reporte {mes}", ["Concepto", "Valor"], [
        ("Ventas netas del mes", r["total"]),
        ("Meta del mes", meta.valor if meta else None),
        ("Cumplimiento de la meta", r["total"] / meta.valor if meta else None),
        ("Devoluciones", r["devoluciones"]),
        ("Abonado a proveedores", sum(p.valor for p in pagos)),
        ("Comisiones liquidadas", sum(l.total for l in liquidaciones)),
        (f"Deuda con proveedores al {hoy:%d/%m/%Y}", sum(f.saldo for f in pendientes)),
        ("  De la cual está vencida", sum(f.saldo for f in pendientes if f.estado == "vencida")),
    ], formatos=pesos(2))
    ws.cell(row=FILA_ENCABEZADO + 3, column=2).number_format = FORMATO_PCT  # fila "Cumplimiento"

    # 2. Ventas diarias
    columnas_valor = list(range(2, len(CANALES) + 4))
    escribir_hoja(wb.create_sheet("Ventas diarias"), f"Ventas diarias — {mes}",
                  ["Fecha", *CANALES, "Devoluciones", "Neto"],
                  [(fecha, *[d["canales"].get(c, 0) for c in CANALES], d["devoluciones"], d["neto"])
                   for fecha, d in r["por_dia"].items()],
                  formatos={1: FORMATO_FECHA, **pesos(*columnas_valor)}, totales=columnas_valor)

    # 3. Por vendedora
    escribir_hoja(wb.create_sheet("Por vendedora"), f"Ventas por vendedora — {mes}",
                  ["Vendedora", "Ventas netas"], list(r["por_vendedora"].items()),
                  formatos=pesos(2), totales=[2])

    # 4. Deudas (foto al día de hoy)
    escribir_hoja(wb.create_sheet("Deudas"), f"Deudas con proveedores al {hoy:%d/%m/%Y}",
                  ["Proveedor", "Factura", "Emisión", "Fecha límite", "Total", "Abonado", "Saldo", "Estado", "Días vencida"],
                  [(f.proveedor.nombre, f.numero, f.fecha_emision, f.fecha_limite, f.total, f.abonado, f.saldo,
                    f.estado, max(0, -f.dias_para_vencer)) for f in pendientes],
                  formatos={3: FORMATO_FECHA, 4: FORMATO_FECHA, **pesos(5, 6, 7)}, totales=[5, 6, 7])

    # 5. Abonos del mes
    escribir_hoja(wb.create_sheet("Abonos"), f"Abonos a proveedores — {mes}",
                  ["Fecha", "Proveedor", "Valor", "Medio", "Referencia", "Aplicado a"],
                  [(p.fecha, p.proveedor.nombre, p.valor, p.medio, p.referencia or "",
                    ", ".join(a.factura.numero for a in p.aplicaciones)) for p in pagos],
                  formatos={1: FORMATO_FECHA, **pesos(3)}, totales=[3])

    # 6. Comisiones del mes
    escribir_hoja(wb.create_sheet("Comisiones"), f"Comisiones — {mes}",
                  ["Quincena", "Vendedora", "Ventas netas", "Tasa", "Comisión", "Bono", "Total", "Estado", "Pagada el", "Soporte"],
                  [(etiqueta_quincena(l.inicio, l.fin), l.vendedora.nombre, l.ventas_netas, l.tasa_bp / 10000,
                    l.comision, l.bono, l.total, l.estado, l.fecha_pago, "Sí" if l.tiene_soporte else "No")
                   for l in liquidaciones],
                  formatos={4: "0.00%", 9: FORMATO_FECHA, **pesos(3, 5, 6, 7)}, totales=[5, 6, 7])

    archivo = io.BytesIO()
    wb.save(archivo)
    archivo.seek(0)
    return archivo