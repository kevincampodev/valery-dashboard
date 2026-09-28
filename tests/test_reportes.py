import io
from datetime import date

from openpyxl import Workbook, load_workbook

from app.services.reportes import escribir_hoja, FORMATO_PESOS, FORMATO_FECHA


def test_escribir_hoja_con_formatos_y_totales():
    wb = Workbook()
    escribir_hoja(wb.active, "Prueba", ["Fecha", "Valor"],
                  [(date(2026, 9, 1), 1000), (date(2026, 9, 2), 2500)],
                  formatos={1: FORMATO_FECHA, 2: FORMATO_PESOS}, totales=[2])

    archivo = io.BytesIO()
    wb.save(archivo)
    archivo.seek(0)
    ws = load_workbook(archivo).active   # releer confirma que el archivo es un .xlsx válido

    assert ws["A1"].value == "Prueba"
    assert ws["A3"].value == "Fecha"
    assert ws["B4"].value == 1000
    assert ws["B4"].number_format == FORMATO_PESOS
    assert ws["B6"].value == "=SUM(B4:B5)"
    assert ws.freeze_panes == "A4"