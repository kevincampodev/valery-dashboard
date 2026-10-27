"""
Lee el Excel de facturas del almacén (una factura puede ocupar varias filas).

Formato de las columnas (fila 1 = encabezados):
  A # factura | B fecha | C nombre cliente | D # telefono | E ref | F descripcion | G cantidad | H color
  I valor unitario | J Subtotal | K valor total factura | L metodo de pago | M Valor metodo de pago
  N Valor pagado | O Valor que debe | P vendedora | Q de donde viene?

La fila con número en la columna A abre una factura. Las filas de abajo sin número son más prendas
o más pagos de esa misma factura.
"""
import re
from datetime import date, datetime

from openpyxl import load_workbook

from .clientes import normalizar_telefono

VENDEDORA_MAYORISTA = "Al por mayor"          # nombre en el panel para las ventas "AL X MAYOR"

# Nombres distintos en el Excel para la misma vendedora del panel
ALIAS_VENDEDORAS = {"Alejandra Diaz": "Alejandra", "Alejandra Díaz": "Alejandra"}

# Secuencia de referencias del catálogo (primeros 3 dígitos) -> descripción
DESCRIPCION_POR_PREFIJO = {
    "050": "Vestido Corto", "051": "Vestido Corto", "057": "Vestido Largo",
    "070": "Enterizo", "071": "Enterizo",
    "090": "Conjunto Pantalon y Blusa", "091": "Conjunto Pantalon y Blusa", "092": "Conjunto Pantalon y Blusa",
    "093": "Conjunto Falda y Blusa",
}
SIN_DESCRIPCION = "Sin descripción"

MEDIOS = {"efectivo": "Efectivo", "bancolombia": "Bancolombia", "nequi": "Nequi", "bold qr": "Bold QR",
          "bold datafono": "Bold Datáfono", "bold datáfono": "Bold Datáfono", "addi": "Addi",
          "transferencia": "Transferencia", "datafono": "Datáfono", "datáfono": "Datáfono"}

CONTACTOS = {"organico": "Orgánico", "orgánico": "Orgánico", "tiktok": "TikTok", "tik tok": "TikTok",
             "facebook": "Facebook", "instagram": "Instagram", "voz a voz": "Voz a voz",
             "cliente": "Ya era cliente"}

# Posición de cada columna (0 = A)
NUM, FECHA, NOMBRE, TEL, REF, DESC, CANT, COLOR, PRECIO, SUBTOTAL, TOTAL, MEDIO, VALOR_MEDIO, \
    PAGADO, DEBE, VENDEDORA, CONTACTO = range(17)


def _texto(valor):
    return " ".join(str(valor).split()) if valor is not None else ""


def _numero(valor):
    if isinstance(valor, (int, float)):
        return int(round(valor))
    digitos = re.sub(r"\D", "", str(valor or ""))
    return int(digitos) if digitos else 0


def leer_fecha(valor):
    """Fecha de la factura. Corrige el año mal digitado de 3 cifras: '03/10/226' -> 3 de octubre de 2026."""
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    partes = re.fullmatch(r"(\d{1,2})/(\d{1,2})/(\d{2,4})", _texto(valor))
    if not partes:
        return None
    dia, mes, anio = (int(p) for p in partes.groups())
    if anio < 1000:
        anio = 2000 + anio % 100
    try:
        return date(anio, mes, dia)
    except ValueError:
        return None


def normalizar_medio(texto):
    return MEDIOS.get(_texto(texto).lower(), "Otro")


def normalizar_contacto(texto):
    clave = _texto(texto).lower()
    if not clave:
        return None                                   # en el panel se ve como "Sin dato"
    return CONTACTOS.get(clave, "Otro")


def normalizar_vendedora(texto):
    """'andrea' -> 'Andrea'. 'AL X MAYOR' -> vendedora general de mayoristas."""
    limpio = _texto(texto)
    if not limpio:
        return None
    if limpio.upper().replace(".", "") in ("AL X MAYOR", "AL POR MAYOR"):
        return VENDEDORA_MAYORISTA
    nombre = limpio.title()
    return ALIAS_VENDEDORAS.get(nombre, nombre)


def inferir_descripcion(referencia, precio, conocidas):
    """
    Descripción para una prenda que llegó sin ella. Reglas acordadas con la administración, en orden:
    1. La misma referencia ya tiene una sola descripción en otra factura.
    2. Secuencia del catálogo: 050/051 corto, 057 largo, 070/071 enterizo, 090-092 conjunto, 093 falda.
    3. "Zapatos", o 4 dígitos que empiezan por 1 o 2 y valen hasta $150.000 -> Calzado.
    4. 1 a 3 dígitos, o 5 dígitos que empiezan por 18, 25 o 26, de $200.000 o más -> Vestido de Gala.
    5. Solo dice "Vestido" -> Vestido (no se sabe si es corto, largo o de gala).
    """
    ref = referencia.strip()
    clave = ref.lower()
    opciones = conocidas.get(clave, set())
    if len(opciones) == 1:
        return next(iter(opciones))
    if re.fullmatch(r"\d{5}", ref) and ref[:3] in DESCRIPCION_POR_PREFIJO:
        return DESCRIPCION_POR_PREFIJO[ref[:3]]
    if clave == "zapatos" or (re.fullmatch(r"[12]\d{3}", ref) and precio <= 150_000):
        return "Calzado"
    if (re.fullmatch(r"\d{1,3}", ref) or re.fullmatch(r"(18|25|26)\d{3}", ref)) and precio >= 200_000:
        return "Vestido de Gala"
    if clave == "vestido":
        return "Vestido"
    if re.search(r"[a-záéíóúñ]", clave) and clave != "zapatos":
        return ref                      # la referencia ya es una descripción ("Vestido de Gala", "Chanel Satin")
    return SIN_DESCRIPCION


def completar_descripciones(ventas):
    """Llena las descripciones vacías. Devuelve (cuántas se completaron, lista de las que no se pudo)."""
    conocidas = {}
    for venta in ventas:
        for linea in venta["lineas"]:
            if linea["descripcion"]:
                conocidas.setdefault(linea["referencia"].lower(), set()).add(linea["descripcion"])
    completadas, pendientes = 0, []
    for venta in ventas:
        for linea in venta["lineas"]:
            if linea["descripcion"]:
                continue
            linea["descripcion"] = inferir_descripcion(linea["referencia"], linea["precio_unitario"], conocidas)
            if linea["descripcion"] == SIN_DESCRIPCION:
                pendientes.append((venta["numero"], linea["referencia"], linea["precio_unitario"]))
            else:
                completadas += 1
    return completadas, pendientes


SIN_NOMBRE = {"nn", "n.n", "n/n", "na", "n/a", "sin nombre", "cliente", "x"}


def normalizar_cliente(texto):
    """'marly vannesa' -> 'Marly Vannesa'. 'NN' o 'N/A' -> None (sin nombre)."""
    nombre = _texto(texto)
    if nombre.lower().replace(" ", "") in {n.replace(" ", "") for n in SIN_NOMBRE}:
        return None
    return nombre.title()[:120] or None


def _es_abono(fila):
    return "ABONO" in _texto(fila[REF]).upper()


def _agrupar(filas):
    """Junta cada factura con sus filas de continuación."""
    facturas, actual = [], None
    for numero_fila, fila in filas:
        if fila[NUM] is not None:
            actual = {"fila": numero_fila, "cabecera": fila, "filas": []}
            facturas.append(actual)
        if actual is not None:
            actual["filas"].append(fila)
    return facturas


def leer_facturas(ruta):
    """Devuelve (ventas, abonos, avisos). No toca la base de datos."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        hoja = wb.worksheets[0]
        filas = []
        for numero_fila, fila in enumerate(hoja.iter_rows(min_row=2, values_only=True), start=2):
            fila = tuple(fila) + (None,) * 17
            if any(v is not None for v in fila[:17]):
                filas.append((numero_fila, fila[:17]))
    finally:
        wb.close()

    ventas, abonos, avisos = [], [], []
    for grupo in _agrupar(filas):
        cab = grupo["cabecera"]
        numero = _numero(cab[NUM])
        etiqueta = f"Factura #{numero}"
        fecha = leer_fecha(cab[FECHA])
        if not fecha:
            avisos.append(f"{etiqueta}: fecha inválida ({cab[FECHA]!r}), no se importa.")
            continue
        if not isinstance(cab[FECHA], (date, datetime)):
            avisos.append(f"{etiqueta}: fecha '{cab[FECHA]}' corregida a {fecha:%d/%m/%Y}.")

        # Pagos (pueden venir en varias filas). Si se repite el medio, se suma.
        pagos = {}
        for fila in grupo["filas"]:
            valor = _numero(fila[VALOR_MEDIO])
            if fila[MEDIO] and valor > 0:
                medio = normalizar_medio(fila[MEDIO])
                pagos[medio] = pagos.get(medio, 0) + valor

        if _es_abono(cab):
            valor = _numero(cab[TOTAL]) or sum(pagos.values())
            medio = next(iter(pagos), "Otro")
            abonos.append({"numero": numero, "fecha": fecha, "valor": valor, "medio": medio,
                           "cliente": _texto(cab[NOMBRE]) or None})
            continue

        # Prendas: filas con precio. Las filas que solo traen un pago se saltan.
        lineas = []
        for fila in grupo["filas"]:
            precio = _numero(fila[PRECIO])
            if precio <= 0:
                continue
            referencia, descripcion = _texto(fila[REF]), _texto(fila[DESC])
            lineas.append({
                "referencia": (referencia or descripcion or "Sin referencia")[:30],
                "descripcion": descripcion[:80] or None,
                "cantidad": _numero(fila[CANT]) or 1,
                "precio_unitario": precio,
                "color": _texto(fila[COLOR]).capitalize()[:30] or None,
            })
        if not lineas:
            avisos.append(f"{etiqueta}: no tiene prendas con precio, no se importa.")
            continue

        total = sum(l["cantidad"] * l["precio_unitario"] for l in lineas)
        total_excel = _numero(cab[TOTAL])
        if (total_excel and total_excel != total and len(lineas) == 1
                and total_excel % lineas[0]["cantidad"] == 0):
            # Una sola prenda con el subtotal digitado distinto al precio: manda lo que se cobró.
            anterior = lineas[0]["precio_unitario"]
            lineas[0]["precio_unitario"] = total_excel // lineas[0]["cantidad"]
            total = total_excel
            avisos.append(f"{etiqueta}: precio unitario {anterior} no cuadra con el subtotal; "
                          f"se toma {lineas[0]['precio_unitario']}.")
        elif total_excel and total_excel != total:
            avisos.append(f"{etiqueta}: el total del Excel ({total_excel}) no cuadra con las prendas ({total}). "
                          f"Se usa el de las prendas.")

        # Decisión de la administración: lo que quedó debiendo se toma como pagado.
        pagado = sum(pagos.values())
        if not pagos:
            pagos = {"Otro": total}
            avisos.append(f"{etiqueta}: sin método de pago, queda como 'Otro'.")
        elif pagado < total:
            primero = next(iter(pagos))
            pagos[primero] += total - pagado
            avisos.append(f"{etiqueta}: debía {total - pagado}, se marca como pagada ({primero}).")
        elif pagado > total:
            avisos.append(f"{etiqueta}: los pagos ({pagado}) suman más que la factura ({total}). Se guardan tal cual.")

        telefono_texto = _texto(cab[TEL])
        telefono = normalizar_telefono(telefono_texto)
        if telefono_texto and not telefono:
            avisos.append(f"{etiqueta}: teléfono '{telefono_texto}' no válido, el cliente queda sin teléfono.")

        vendedora = normalizar_vendedora(cab[VENDEDORA])
        ventas.append({
            "numero": numero,
            "fecha": fecha,
            "canal": "Mayorista" if vendedora == VENDEDORA_MAYORISTA else "Minorista",
            "vendedora": vendedora,
            "cliente": normalizar_cliente(cab[NOMBRE]),
            "telefono": telefono,
            "medio_contacto": normalizar_contacto(cab[CONTACTO]),
            "valor": total,
            "prendas": sum(l["cantidad"] for l in lineas),
            "pagos": pagos,
            "lineas": lineas,
        })
    completadas, pendientes = completar_descripciones(ventas)
    if completadas:
        avisos.append(f"{completadas} prenda(s) sin descripción se completaron con las reglas de referencia.")
    for numero, referencia, precio in pendientes:
        avisos.append(f"Factura #{numero}: la referencia '{referencia}' ({precio}) quedó '{SIN_DESCRIPCION}'.")
    return ventas, abonos, avisos


def numeros_faltantes(ventas, abonos):
    """Números de factura que no aparecen entre el primero y el último."""
    numeros = {v["numero"] for v in ventas} | {a["numero"] for a in abonos}
    if not numeros:
        return []
    return sorted(set(range(min(numeros), max(numeros) + 1)) - numeros)
