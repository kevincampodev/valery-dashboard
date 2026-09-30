from datetime import date, datetime

from openpyxl import load_workbook

ENCABEZADOS = ["FILA ORIGINAL", "FECHA", "VENDEDOR", "VALOR", "MEDIO DE CONTACTO", "CANAL"]
CONTACTOS = {"organico": "Orgánico", "orgánico": "Orgánico", "tik tok": "TikTok", "tiktok": "TikTok",
             "voz a voz": "Voz a voz"}


def normalizar_vendedor(texto):
    limpio = " ".join(str(texto or "").split())
    return limpio.title() or None


def normalizar_contacto(texto):
    clave = " ".join(str(texto or "").lower().split())
    if not clave:
        return None
    return CONTACTOS.get(clave, "Otro")


def normalizar_canal(texto):
    # Decisión de la administración: lo que no esté marcado Mayorista (incluido REVISAR) es minorista
    return "Mayorista" if str(texto or "").strip().lower() == "mayorista" else "Minorista"


def a_fecha(valor):
    if isinstance(valor, datetime):
        return valor.date()
    if isinstance(valor, date):
        return valor
    texto = str(valor or "").strip()
    for formato in ("%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(texto, formato).date()
        except ValueError:
            continue
    return None


def leer_revision(ruta):
    """Lee la hoja 'Ventas' del archivo de revisión. Devuelve (ventas, omitidas)."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if "Ventas" not in wb.sheetnames:
            raise ValueError("El archivo no tiene una hoja llamada 'Ventas'.")
        filas = wb["Ventas"].iter_rows(values_only=True)
        encabezado = [str(c or "").strip().upper() for c in next(filas)]
        faltan = [e for e in ENCABEZADOS if e not in encabezado]
        if faltan:
            raise ValueError(f"Faltan columnas en la hoja 'Ventas': {', '.join(faltan)}")
        col = {e: encabezado.index(e) for e in ENCABEZADOS}

        ventas, omitidas = [], []
        for numero, fila in enumerate(filas, start=2):
            fila = tuple(fila) + (None,) * len(encabezado)
            if not any(fila):
                continue
            referencia = fila[col["FILA ORIGINAL"]] or f"fila {numero}"
            fecha = a_fecha(fila[col["FECHA"]])
            valor = fila[col["VALOR"]]
            if not fecha:
                omitidas.append((referencia, "fecha inválida"))
                continue
            if not isinstance(valor, (int, float)) or valor <= 0:
                omitidas.append((referencia, "sin valor"))
                continue
            ventas.append({
                "fila": referencia,
                "fecha": fecha,
                "vendedora": normalizar_vendedor(fila[col["VENDEDOR"]]),
                "valor": int(round(valor)),
                "canal": normalizar_canal(fila[col["CANAL"]]),
                "medio_contacto": normalizar_contacto(fila[col["MEDIO DE CONTACTO"]]),
            })
        return ventas, omitidas
    finally:
        wb.close()


def resumir(ventas):
    grupos = {"Por mes": {}, "Por canal": {}, "Por medio de contacto": {}, "Por vendedora": {}}
    for v in ventas:
        claves = {
            "Por mes": v["fecha"].strftime("%Y-%m"),
            "Por canal": v["canal"],
            "Por medio de contacto": v["medio_contacto"] or "Sin dato",
            "Por vendedora": v["vendedora"] or "Sin vendedora",
        }
        for grupo, clave in claves.items():
            grupos[grupo][clave] = grupos[grupo].get(clave, 0) + v["valor"]
    grupos["Por mes"] = dict(sorted(grupos["Por mes"].items()))
    for g in ("Por canal", "Por medio de contacto", "Por vendedora"):
        grupos[g] = dict(sorted(grupos[g].items(), key=lambda x: -x[1]))
    return grupos